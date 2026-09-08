import sys
import json
import openpyxl
import copy
import datetime
import tempfile
import os
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def format_date_id(date_str):
    months_id = {
        "01": "Jan", "02": "Feb", "03": "Mar", "04": "Apr", "05": "Mei", "06": "Jun",
        "07": "Jul", "08": "Agu", "09": "Sep", "10": "Okt", "11": "Nov", "12": "Des"
    }
    try:
        parts = date_str.split("-")
        day = str(int(parts[2]))
        month = months_id[parts[1]]
        year = parts[0][2:]
        return f"{day} {month} {year}"
    except:
        return date_str

def parse_time(time_str):
    try:
        parts = time_str.split(":")
        return datetime.time(int(parts[0]), int(parts[1]))
    except:
        return time_str

def is_dosen_luar(row):
    # Sumber kebenaran: field "status" (dari kolom status_dosen di database, diatur
    # lewat halaman List Dosen). Daftar kata kunci di bawah cuma fallback untuk baris
    # lama/offline yang belum membawa field status sama sekali.
    status = row.get("status")
    if status:
        return status == "luar"

    name_lower = row.get("dsn", "").lower()
    luar_keywords = [
        "anggi sandika",
        "adityana",
        "amelia",
        "hendy satria",
        "handy satria",
        "i'oh",
        "iin nurhayati",
        "sonny feisal",
        "sussylawati",
        "tarma"
    ]
    return any(keyword in name_lower for keyword in luar_keywords)

def main():
    try:
        # Baca JSON payload dari stdin
        payload = json.loads(sys.stdin.read())
        
        # Path dihitung relatif terhadap lokasi skrip ini (src/lib/generate_excel.py),
        # supaya tidak bergantung pada drive/folder proyek di komputer tertentu.
        project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        template_path = os.path.join(project_root, "File Pendukung", "Honjar Mei Tahun 2026 PRODI TLM D4.xlsx")
        if not os.path.exists(template_path):
            print(f"ERROR: Template file not found at {template_path}", file=sys.stderr)
            sys.exit(1)
            
        wb = openpyxl.load_workbook(template_path)
        
        month_name = payload.get("monthName", "Mei 2026").upper()
        current_date_str = payload.get("currentDate", "8 Juni 2026")
        rows_data = payload.get("rows", [])
        
        # Kelompokkan ke Dosen Tetap dan Dosen Luar
        rows_tetap = []
        rows_luar = []
        
        for r in rows_data:
            if is_dosen_luar(r):
                rows_luar.append(r)
            else:
                rows_tetap.append(r)
                
        # Urutkan: Dosen -> Tanggal -> Jam Mulai
        rows_tetap.sort(key=lambda x: (x["dsn"], x["tgl"], x["a"]))
        rows_luar.sort(key=lambda x: (x["dsn"], x["tgl"], x["a"]))
        
        # Eksekusi untuk masing-masing Sheet
        for sheet_name, data in [("Tetap", rows_tetap), ("Luar", rows_luar)]:
            sheet = wb[sheet_name]
            
            # Ganti periode pada judul di baris 3
            sheet.cell(row=3, column=1).value = f"PERIODE {month_name}"
            
            # Ubah jenis font untuk semua cell di baris 1 s/d 6 menjadi Times New Roman
            for r_idx in range(1, 7):
                for col_idx in range(1, 16):
                    cell = sheet.cell(row=r_idx, column=col_idx)
                    if cell.font:
                        f_copy = copy.copy(cell.font)
                        f_copy.name = "Times New Roman"
                        cell.font = f_copy
            
            # Ambil style formatting dasar dari baris data ke-7 template
            template_styles = []
            for col in range(1, 12):
                cell = sheet.cell(row=7, column=col)
                template_styles.append({
                    'font': copy.copy(cell.font),
                    'alignment': copy.copy(cell.alignment),
                    'fill': copy.copy(cell.fill),
                    'border': copy.copy(cell.border),
                    'number_format': cell.number_format
                })

            # Tinggi baris data normal, diambil dari baris pertama yang tidak dipakai
            # sebagai baris style (baris 8 pada template asli). Dipakai untuk menimpa
            # tinggi baris 7 bawaan template yang sengaja dibuat sangat kecil (0.95px,
            # baris bantu kosong) — kalau tidak ditimpa, baris data pertama ikut kecil.
            normal_row_height = sheet.row_dimensions[8].height if 8 in sheet.row_dimensions else 15

            # Isian baris pemisah kuning antar kelompok dosen
            yellow_fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")

            # Bersihkan seluruh data lama dari baris 7 sampai baris terakhir secara efisien
            sheet.delete_rows(7, sheet.max_row - 6)

            # Tulis data baru mulai baris 7
            current_dsn = None
            no_counter = 0

            start_row = 7
            row_num = start_row
            for r in data:
                if r["dsn"] != current_dsn:
                    # Baris pemisah kuning sebelum kelompok dosen baru (bukan sebelum yang pertama)
                    if current_dsn is not None:
                        sheet.row_dimensions[row_num].height = normal_row_height
                        for col_idx in range(1, 12):
                            cell = sheet.cell(row=row_num, column=col_idx)
                            cell.value = None
                            style = template_styles[col_idx - 1]
                            if style['border']: cell.border = copy.copy(style['border'])
                            cell.fill = yellow_fill
                        row_num += 1

                    no_counter += 1
                    no_val = no_counter
                    current_dsn = r["dsn"]
                else:
                    no_val = None

                # Tulis data serta formula Excel
                cell_vals = {
                    1: no_val,                           # NO
                    2: r["dsn"],                         # DOSEN
                    3: format_date_id(r["tgl"]),         # TANGGAL
                    4: parse_time(r["a"]),               # MULAI
                    5: parse_time(r["b"]),               # SELESAI
                    6: f"=VALUE(K{row_num})",            # DURASI
                    7: f"=IF(F{row_num}>=300,6,IF(F{row_num}>=250,5,IF(F{row_num}>=200,4,IF(F{row_num}>=150,3,IF(F{row_num}>=100,2,IF(F{row_num}>=50,1,0))))))", # JUMLAH JAM
                    8: r["mk"],                          # MATA KULIAH
                    9: r["met"],                         # METODE
                    10: r["kls"],                        # KELAS
                    11: f'=TEXT(E{row_num}-D{row_num},"[mm]")' # Hitung Menit (Helper)
                }

                sheet.row_dimensions[row_num].height = normal_row_height
                for col_idx in range(1, 12):
                    cell = sheet.cell(row=row_num, column=col_idx)
                    cell.value = cell_vals[col_idx]

                    # Salin style template dan ubah font name ke Times New Roman
                    style = template_styles[col_idx - 1]
                    if style['font']:
                        f_copy = copy.copy(style['font'])
                        f_copy.name = "Times New Roman"
                        cell.font = f_copy
                    if style['alignment']: cell.alignment = copy.copy(style['alignment'])
                    if style['fill']: cell.fill = copy.copy(style['fill'])
                    if style['border']: cell.border = copy.copy(style['border'])
                    if style['number_format']: cell.number_format = style['number_format']

                row_num += 1

            # Tulis Tanda Tangan Ka. Prodi di akhir tabel (3 baris setelah data berakhir)
            last_data_row = row_num - 1
            if len(data) == 0:
                last_data_row = 6

            sig_start = last_data_row + 4
            
            sheet.cell(row=sig_start, column=8).value = f"Cimahi, {current_date_str}"
            sheet.cell(row=sig_start, column=8).font = Font(name="Times New Roman", size=10)
            
            sheet.cell(row=sig_start + 1, column=8).value = "Ka. Prodi TLM D4"
            sheet.cell(row=sig_start + 1, column=8).font = Font(name="Times New Roman", size=10, bold=True)
            
            sheet.cell(row=sig_start + 5, column=8).value = "Gina Khairinisa,M.Imun"
            sheet.cell(row=sig_start + 5, column=8).font = Font(name="Times New Roman", size=10, bold=True)
            
            sheet.cell(row=sig_start + 6, column=8).value = "NID. 4121 282 88"
            sheet.cell(row=sig_start + 6, column=8).font = Font(name="Times New Roman", size=10)
            
        # Simpan ke temp file
        fd, temp_file_path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(temp_file_path)
        
        # Cetak output path ke stdout agar dibaca Next.js
        print(temp_file_path)
        
    except Exception as e:
        print(f"ERROR: {str(e)}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
